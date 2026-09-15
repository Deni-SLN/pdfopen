from fastapi import APIRouter, Depends, HTTPException, BackgroundTasks
from sqlalchemy.orm import Session
from pathlib import Path
import json, uuid
from ..database import get_db
from ..deps import get_current_user
from ..models import User, Job, File as FileModel
from ..config import settings

router=APIRouter(prefix="/api/jobs", tags=["jobs"])

def _enqueue(job_id: str):
    # try RQ, fallback to background thread
    try:
        import redis
        from rq import Queue
        from ..workers.processor import process_job
        r=redis.from_url(settings.REDIS_URL)
        q=Queue("sofia", connection=r)
        q.enqueue(process_job, job_id, job_timeout=600)
        return
    except Exception as e:
        pass
    # fallback: run in background thread via processor directly (async)
    import threading
    from ..workers.processor import process_job
    threading.Thread(target=process_job, args=(job_id,), daemon=True).start()

@router.post("")
def create_job(payload: dict, bg: BackgroundTasks, user: User=Depends(get_current_user), db: Session=Depends(get_db)):
    tool=payload.get("tool")
    params=payload.get("params") or {}
    file_id=payload.get("file_id")
    file_ids=payload.get("file_ids") or []
    if not tool: raise HTTPException(400,"tool required")
    # ownership check
    if file_id:
        f=db.query(FileModel).filter(FileModel.id==file_id).first()
        if not f or f.owner_id!=user.id: raise HTTPException(404,"file not found")
    job=Job(owner_id=user.id, tool=tool, status="queued", progress=0, message="Queued", params=json.dumps(params), input_file_id=file_id or (file_ids[0] if file_ids else None))
    # store file_ids in params if not present
    if file_ids and "file_ids" not in params:
        p=json.loads(job.params); p["file_ids"]=file_ids; job.params=json.dumps(p)
    db.add(job); db.commit(); db.refresh(job)
    _enqueue(job.id)
    return job

@router.get("")
def list_jobs(user: User=Depends(get_current_user), db: Session=Depends(get_db)):
    return db.query(Job).filter(Job.owner_id==user.id).order_by(Job.created_at.desc()).limit(100).all()

@router.get("/{jid}")
def get_job(jid:str, user: User=Depends(get_current_user), db: Session=Depends(get_db)):
    j=db.query(Job).filter(Job.id==jid).first()
    if not j or j.owner_id!=user.id: raise HTTPException(404,"Not found")
    return j

@router.post("/{jid}/cancel")
def cancel_job(jid:str, user: User=Depends(get_current_user), db: Session=Depends(get_db)):
    j=db.query(Job).filter(Job.id==jid).first()
    if not j or j.owner_id!=user.id: raise HTTPException(404,"Not found")
    if j.status in ("completed","failed","cancelled"): return j
    j.status="cancelled"; j.message="Cancelled by user"; db.commit()
    return j

@router.get("/{jid}/download")
def download_job(jid:str, user: User=Depends(get_current_user), db: Session=Depends(get_db)):
    from fastapi.responses import FileResponse
    j=db.query(Job).filter(Job.id==jid).first()
    if not j or j.owner_id!=user.id: raise HTTPException(404,"Not found")
    if j.status!="completed" or not j.output_file_id: raise HTTPException(400,"Job not completed")
    f=db.query(FileModel).filter(FileModel.id==j.output_file_id).first()
    if not f: raise HTTPException(404,"File not found")
    p=Path(f.path)
    if not p.exists(): raise HTTPException(404,"File missing on disk")
    return FileResponse(str(p), filename=f.original_name, media_type="application/octet-stream")

# preview for pdf->csv
@router.get("/{jid}/preview")
def preview(jid:str, user: User=Depends(get_current_user), db: Session=Depends(get_db)):
    j=db.query(Job).filter(Job.id==jid).first()
    if not j or j.owner_id!=user.id: raise HTTPException(404,"Not found")
    if not j.output_file_id: raise HTTPException(400,"No output")
    f=db.query(FileModel).filter(FileModel.id==j.output_file_id).first()
    import csv
    p=Path(f.path)
    if p.suffix==".xlsx":
        try:
            from openpyxl import load_workbook
            wb=load_workbook(str(p), read_only=True)
            ws=wb.active
            rows=[]
            for i,row in enumerate(ws.iter_rows(values_only=True)):
                if i>=30: break
                rows.append([("" if v is None else str(v)) for v in row])
            wb.close()
            return {"rows":rows, "filename":f.original_name}
        except Exception:
            pass
    if p.suffix==".csv":
        with open(p, encoding="utf-8", errors="ignore") as fh:
            rows=[next(csv.reader([line])) if False else None]
        # simple read
        with open(p, encoding="utf-8", errors="ignore") as fh:
            r=csv.reader(fh)
            rows=[]
            for i,row in enumerate(r):
                if i>=30: break
                rows.append(row)
        return {"rows":rows, "filename":f.original_name}
    return {"filename":f.original_name}
