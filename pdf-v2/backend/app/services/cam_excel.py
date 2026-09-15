"""PDF -> Excel CAM converter.

Khusus file "Detail CAM Prioritisasi" (export jsPDF dengan watermark diagonal).
Pipeline:
  1. Hapus watermark (teks yang di-render dengan matriks rotasi) langsung di
     content stream -> garis tabel & teks konten tidak tersentuh.
  2. Ekstrak tabel per halaman dengan pdfplumber (strategi "lines"), rekonstruksi
     isi sel dari koordinat karakter agar wrap tsPDF tersusun kembali dengan benar.
  3. Gabungkan semua halaman ke SATU sheet dengan header baku 32 kolom.
"""
import io, re
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor

CAM_HEADERS = [
    "No CIF", "No Loan", "Kode Cabang", "Kode Outlet", "Nama Nasabah",
    "Nama Produk", "Nama Produk Simplifikasi", "Kol EOM", "Prioritas",
    "KOL SOM", "KOL EOD", "Predictive KOL", "Sisa Pokok Outstanding",
    "Angsuran", "Tgl Due Date", "Tgl Bayar Pokok", "Start Monitoring Restru",
    "End Monitoring Restru", "Proper", "Developer", "Kriteria Nasabah",
    "Kondisi Agunan", "Permasalahan", "Kelengkapan Dokumen", "PIC Desk Coll KP",
    "PIC Desk Coll KC", "PIC Field Collector", "Status Pembinaan",
    "Tgl Janji Bayar", "Int. Telepon", "Int. Visit", "Load Timestamp",
]
NCOL = len(CAM_HEADERS)

# kolom yang isinya angka/kode/tanggal: buang semua spasi lalu perbaiki pola
_NOSPACE_IDX = {0, 1, 2, 3, 12, 13, 14, 15, 16, 17, 18, 28, 29, 30, 31}
_KOL_IDX = {7, 9, 10, 11}
_PRIORITAS_IDX = 8

# --- watermark strip (content stream surgery) -------------------------------
_BT_ET = re.compile(rb"BT(?:(?!BT|ET).)*?ET", re.S)
_TM = re.compile(rb"([\d.\-]+)\s+([\d.\-]+)\s+([\d.\-]+)\s+([\d.\-]+)\s+([\d.\-]+)\s+([\d.\-]+)\s+Tm")


def _strip_watermark_stream(data: bytes):
    removed = 0
    def repl(m):
        nonlocal removed
        block = m.group(0)
        for tm in _TM.finditer(block):
            b = float(tm.group(2)); c = float(tm.group(3))
            if abs(b) > 1e-4 or abs(c) > 1e-4:  # matriks rotasi = watermark
                removed += 1
                return b""
        return block
    return _BT_ET.sub(repl, data), removed


def remove_cam_watermark_bytes(inp: Path) -> bytes:
    """Kembalikan byte PDF tanpa watermark (teks rotasi dihapus dari stream)."""
    import fitz
    doc = fitz.open(str(inp))
    try:
        for page in doc:
            xrefs = page.get_contents()
            if not xrefs:
                continue
            if len(xrefs) > 1:
                page.clean_contents()
                xrefs = page.get_contents()
            xref = xrefs[0]
            new, n = _strip_watermark_stream(doc.xref_stream(xref))
            if n:
                doc.update_stream(xref, new)
        return doc.tobytes(garbage=4, deflate=True, clean=True)
    finally:
        doc.close()


# --- normalisasi nilai -------------------------------------------------------
def _norm_nospace(s: str) -> str:
    s = re.sub(r"\s+", "", s)
    s = re.sub(r"(\d{4}-\d{2}-\d{2})(\d{2}:\d{2}:\d{2})", r"\1 \2", s)  # tanggal jam
    s = re.sub(r"^Rp", "Rp ", s)                                         # nominal
    return s

def _norm_kol(s: str) -> str:
    s = re.sub(r"\s+", "", s)
    return re.sub(r"^KOL(?=[\d.])", "KOL ", s)                           # KOL2.1 -> KOL 2.1

def _norm_prioritas(s: str) -> str:
    s = re.sub(r"\s+", "", s)
    return re.sub(r"^PRIORITAS(?=\d)", "PRIORITAS ", s)                  # PRIORITAS1 -> PRIORITAS 1

def _norm_col(i: int, c) -> str:
    c = (c or "").strip()
    if not c:
        return ""
    if i in _KOL_IDX:
        return _norm_kol(c)
    if i == _PRIORITAS_IDX:
        return _norm_prioritas(c)
    if i in _NOSPACE_IDX:
        return _norm_nospace(c)
    return re.sub(r"\s{2,}", " ", c)


# --- rekonstruksi sel dari koordinat karakter --------------------------------
def _cell_text(chars, bbox):
    """Susun teks sel; spasi antar-kata dari gap horizontal, dan join antar-baris
    dibedakan antara potongan kata (tight) vs batas kata (pakai spasi)."""
    if not chars:
        return "", False
    x0c, _topc, x1c, _botc = bbox
    pad = max(0.0, min(c["x0"] for c in chars) - x0c)

    # cluster jadi baris berdasarkan posisi vertikal
    chars = sorted(chars, key=lambda c: (round(c["top"], 1), c["x0"]))
    lines, cur, cur_top = [], [], None
    for ch in chars:
        if cur_top is None or abs(ch["top"] - cur_top) <= 2.0:
            cur.append(ch)
            if cur_top is None:
                cur_top = ch["top"]
        else:
            lines.append(cur); cur = [ch]; cur_top = ch["top"]
    if cur:
        lines.append(cur)

    parsed = []  # (text, line_chars, word_chars_list)
    for ln in lines:
        ln = sorted(ln, key=lambda c: c["x0"])
        size = ln[0].get("size", 5) or 5
        thr = max(0.75, size * 0.18)  # gap > thr dianggap spasi
        s, words, cur_w = "", [], [ln[0]]
        prev = None
        for ch in ln:
            if prev is not None:
                if ch["x0"] - prev["x1"] > thr:
                    s += " "
                    words.append(cur_w); cur_w = [ch]
                else:
                    cur_w.append(ch)
            s += ch["text"]
            prev = ch
        words.append(cur_w)
        parsed.append((s.strip(), ln, words))

    out = parsed[0][0]
    U = x1c - pad            # batas kanan area teks
    W = (x1c - x0c) - 2 * pad  # lebar wrap yang dipakai generator PDF
    for i in range(1, len(parsed)):
        _ps, prev_ln, prev_words = parsed[i - 1]
        cur_s, cur_ln, cur_words = parsed[i]
        if not prev_words or not cur_words or not prev_ln or not cur_ln:
            out += (" " + cur_s) if cur_s else ""
            continue
        lw, fw = prev_words[-1], cur_words[0]
        w1 = lw[-1]["x1"] - lw[0]["x0"]      # lebar kata terakhir baris atas
        w2 = fw[-1]["x1"] - fw[0]["x0"]      # lebar kata pertama baris bawah
        gap_right = U - prev_ln[-1]["x1"]    # sisa ruang setelah karakter terakhir
        nch_w = cur_ln[0]["x1"] - cur_ln[0]["x0"]
        # potongan kata (chunk): kata gabungan melebihi lebar baris DAN baris atas
        # terisi penuh sampai tidak muat 1 karakter lagi
        if (w1 + w2) > W and gap_right < nch_w:
            out += cur_s
        else:
            out += " " + cur_s
    # flag "filled": baris terakhir terisi penuh sampai tepi kanan -> kata
    # kemungkinan berlanjut ke baris tabel berikutnya (wrap antar-row)
    filled = False
    last = parsed[-1]
    if last and last[1]:
        last_ln = last[1]
        gap_right = U - last_ln[-1]["x1"]
        ws = [c["x1"] - c["x0"] for c in last_ln if c["x1"] > c["x0"]] or [1.0]
        avg_w = sum(ws) / len(ws)
        filled = gap_right < avg_w * 0.6
    return re.sub(r"\s{2,}", " ", out).strip(), filled


_TSET = {"vertical_strategy": "lines", "horizontal_strategy": "lines",
         "intersection_tolerance": 5, "snap_tolerance": 3, "join_tolerance": 3}


def _extract_page(page):
    import bisect
    tables = page.find_tables(_TSET)
    out = []
    for tbl in tables:
        bbox = tbl.bbox
        inside = [c for c in page.chars
                  if bbox[0] - 1 <= (c["x0"] + c["x1"]) / 2 <= bbox[2] + 1
                  and bbox[1] - 1 <= (c["top"] + c["bottom"]) / 2 <= bbox[3] + 1]
        centers = sorted((((c["x0"] + c["x1"]) / 2, c) for c in inside), key=lambda t: t[0])
        xs = [t[0] for t in centers]
        for row in tbl.rows:
            r = []
            cont = set()
            for j, cell in enumerate(row.cells):
                if cell is None:
                    r.append(""); continue
                lo = bisect.bisect_left(xs, cell[0] - 0.01)
                hi = bisect.bisect_right(xs, cell[2] - 0.01)
                cc = [c for x, c in centers[lo:hi]
                      if cell[1] <= (c["top"] + c["bottom"]) / 2 < cell[3]]
                txt, filled = _cell_text(cc, cell)
                if filled and txt:
                    cont.add(j)
                r.append(txt)
            out.append((r, cont))
    return out


# --- xlsx writer --------------------------------------------------------------
def _write_xlsx(out_path: Path, data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "CAM"
    ws.append(CAM_HEADERS)
    hf = PatternFill("solid", fgColor="1E3A8A")
    for j in range(1, NCOL + 1):
        c = ws.cell(row=1, column=j)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = hf
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in data:
        row = list(row[:NCOL]) + [""] * (NCOL - len(row))
        ws.append([(c if c is not None else "") for c in row])
    widths = [len(h) for h in CAM_HEADERS]
    for row in data:
        for i, v in enumerate(row[:NCOL]):
            l = len(v or "")
            if l > widths[i]:
                widths[i] = min(l, 45)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(w + 2, 45)
    ws.freeze_panes = "A2"
    wb.save(str(out_path))


# --- entry point ---------------------------------------------------------------
def pdf_to_excel_cam(inp: Path, out_path: Path, progress_cb=None):
    """Konversi PDF CAM -> 1 file XLSX (semua halaman jadi 1 sheet)."""
    buf = remove_cam_watermark_bytes(inp)
    if progress_cb:
        progress_cb(20, "Watermark dihapus ✓ • Ekstraksi tabel")
    import pdfplumber
    pdf = pdfplumber.open(io.BytesIO(buf))
    try:
        n = len(pdf.pages)
        page_rows = [None] * n
        def work(args):
            i, page = args
            try:
                return i, _extract_page(page)
            except Exception:
                return i, []
        with ThreadPoolExecutor(max_workers=4) as ex:
            for i, rows in ex.map(work, list(enumerate(pdf.pages))):
                page_rows[i] = rows
                if progress_cb and (i % 20 == 0 or i == n - 1):
                    progress_cb(20 + int(55 * (i + 1) / max(n, 1)),
                                f"Ekstraksi halaman {i + 1}/{n}")
    finally:
        pdf.close()

    data, conts_all = [], []
    for rows in page_rows:
        if not rows:
            continue
        for row, cont in rows:
            if not any((c or "").strip() for c in row):
                continue
            joined = " ".join((c or "") for c in row[:3])
            if "No CIF" in joined:  # header berulang tiap halaman -> skip
                continue
            norm = [_norm_col(i, c) for i, c in enumerate(row)]
            # baris pecahan (wrap tinggi): tanpa No CIF & No Loan -> gabung ke baris sebelumnya.
            # tight-join hanya utk kolom yang baris terakhirnya penuh sampai tepi (kelanjutan kata)
            if data and not norm[0].strip() and not norm[1].strip():
                prev, prev_cont = data[-1], conts_all[-1]
                for i in range(NCOL):
                    v = norm[i]
                    if not v:
                        continue
                    sep = "" if (i in prev_cont or i in _KOL_IDX or i in _NOSPACE_IDX) else " "
                    prev[i] = _norm_col(i, prev[i] + sep + v)
                continue
            data.append(norm)
            conts_all.append(cont)
    if not data:
        raise ValueError("Tidak ada baris data CAM yang bisa diekstrak. Pastikan file adalah Detail CAM Prioritisasi.")
    if progress_cb:
        progress_cb(80, "Menulis file Excel")
    _write_xlsx(out_path, data)
    if progress_cb:
        progress_cb(90, f"Selesai • {len(data)} baris data")
    return len(data)